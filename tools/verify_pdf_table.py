#!/usr/bin/env python3
r"""
判"表格印成 PDF"那一步：排片单由 openpyxl 读原件生成，`:core` 那份排版器排成页，桌面 PDFBox 画出来，
再用 **pdfplumber 的词与线坐标**、**MuPDF(fitz)** 与 **poppler(pdftotext)** 三套独立读路把 PDF 读回来核位置。

跑法（先落夹具与产物）：
    python tools/make_pdf_table_fixtures.py
    GRADLE_USER_HOME=D:/gradle-home ./gradlew :core:jar
    for f in book wide long; do PDFPROBE_OUT=D:/fflogs/pdfprobe bash tools/pdfprobe/run.sh pdftable \
        C:/Windows/Fonts/simhei.ttf core/build/pdftable/$f.plan.tsv core/build/pdftable/$f.pdf; done
    for f in book wide long; do PDFPROBE_OUT=D:/fflogs/pdfprobe bash tools/pdfprobe/run.sh xlsxpdf \
        C:/Windows/Fonts/simhei.ttf core/build/pdftable/$f.xlsx core/build/pdftable/$f-app.pdf; done
    python tools/verify_pdf_table.py

为什么排片单不能让 Kotlin 那边生成：格子内容若由我们自己的 xlsx 读法说了算，PDF 排错的列一旦与
读错的列对上，判据就成了自己圆自己 —— 所以"文件里有什么"只问 openpyxl，"纸上是哪一格"只问
pdfplumber 的坐标，两边都不沾我们自己的读法。

十五条判据：
  1  六份 PDF（每样两份）与说明都在且非空；60 行那份真的翻了页
  2  纸上按先后读出来的每一带，与排片单的行序对得上（跨页多出来的表头带允许，别的都不许）
  3  每一带逐格对：格子的文字落在它自己那一列的 x 区间里（列序、内容一起对）
  4  网格真画出来了：纸上每一带都归得到排片单里的一张表、列数与那张表一致；框上下闭合，
     每条横线都贯穿这段表
  5  表头下面那条线比别的重；没写表头的表不许有重线
  6  跨页那份：第二页第一带就是表头那一行
  7  压窄的说明诚实：notes 说"第 N 列压到 Xpt"，纸上量出来就在 Xpt 上下
  8  格子里的硬换行写在两条基线上，不是挤成一行
  9  每张表上面那行小标题是表名、排在它自己那张表上面，且字号比正文大
 10  一整张真空的表不占一节；表外那句话排在表后面
 11  MuPDF 与 poppler（两套独立解析）读回来的字数都与原件相当，且两家数目一致
     注：poppler 必须 `-enc UTF-8`。这台机器上默认按 Latin-1 出，中文**全部丢掉**（0 个字），
     差点被我写成"我们的 PDF 缺中文文字层"这条产品限制 —— 是调用姿势错，加了这个参数后
     pdftotext / pdfminer / MuPDF 三家从同一份文件里读到一样的 158 个中文
 12  没有一个字越出页边距；压窄也不许压成零宽（最窄的列还装得下一个字）
 13  没有空白页（排了几页就几页有东西）
 14  从真 .xlsx 走**应用那条链子**（我们的 xlsx 读法 → PrintPlanner → PDF）：纸上逐带与排片单一致，
     空表不占一节，也不重复表头（普通区域没声明表头）
 15  同一条链子的说明诚实：notes 里"几张表 · 几行"、每张几行几列，与 openpyxl 数出来的对得上
"""
from __future__ import annotations

import os
import re
import subprocess
import sys

import openpyxl
import pdfplumber

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from verify_odt_read import check, results  # noqa: E402

BUILD = os.path.join("core", "build", "pdftable")
PAGE_WIDTH = 595.0
MARGIN = 56.0
FIXTURES = ("book", "wide", "long")
# 每个夹具两份 PDF：一份由排片单（openpyxl 读的）驱动，一份由**应用那条链子**（我们的 xlsx 读法）驱动
ALL_PRODUCTS = FIXTURES + tuple(name + "-app" for name in FIXTURES)


def flat(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def parse_plan(name: str) -> list:
    """排片单 → 按顺序的条目：("表", 表名, 有无表头, 行) / ("文", 一句) / ("空", 表名)。"""
    out = []
    with open(os.path.join(BUILD, "%s.plan.tsv" % name), encoding="utf-8") as handle:
        lines = handle.read().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("#SHEET\t"):
            head = line.split("\t")
            rows = []
            index += 1
            while index < len(lines) and not lines[index].startswith("#"):
                rows.append([flat(cell.replace("\\n", "\n")) for cell in lines[index].split("\t")])
                index += 1
            out.append(("表", head[1], head[2].endswith("=1"), rows))
            continue
        if line.startswith("#TEXT\t"):
            out.append(("文", line[6:]))
        elif line.startswith("#EMPTY\t"):
            out.append(("空", line[7:]))
        index += 1
    return out


OPENED = []
PAGES = {}


def pages_of(name: str) -> list:
    """页对象是惰性的（要到用时才读流），所以文档句柄得一直开着，进程退出再散。"""
    if name in PAGES:
        return PAGES[name]
    handle = pdfplumber.open(os.path.join(BUILD, "%s.pdf" % name))
    OPENED.append(handle)
    PAGES[name] = list(handle.pages)
    return PAGES[name]


def words_on(page):
    """extract_words 默认不返回字号，要额外要 —— 判字号的那几条靠它。"""
    return page.extract_words(extra_attrs=["size"])


def segments(page) -> list:
    """这一页上的每一段表：竖线按 (顶, 底) 归组得列边界，横线给出行带。"""
    verticals = [line for line in page.lines if abs(line["x0"] - line["x1"]) < 0.6]
    horizontals = [line for line in page.lines if abs(line["top"] - line["bottom"]) < 0.6]
    groups = {}
    for line in verticals:
        groups.setdefault((round(line["top"], 1), round(line["bottom"], 1)), set()).add(round(line["x0"], 1))
    out = []
    for (top, bottom), xs in groups.items():
        edges = sorted(xs)
        rows = sorted({round(line["top"], 1) for line in horizontals
                       if top - 1.0 <= line["top"] <= bottom + 1.0 and line["x0"] <= edges[0] + 1.0})
        heavy = [round(line["top"], 1) for line in horizontals
                 if line["linewidth"] > 1.0 and top - 1.0 <= line["top"] <= bottom + 1.0]
        out.append({"top": top, "bottom": bottom, "edges": edges, "rows": rows, "heavy": heavy})
    return sorted(out, key=lambda item: item["top"])


def bands(page, segment) -> list:
    """一段表读回来：每一带 = 那一横条里各列区间的字按 x 连起来的串。"""
    words = words_on(page)
    out = []
    for index in range(len(segment["rows"]) - 1):
        top, bottom = segment["rows"][index], segment["rows"][index + 1]
        cells = []
        for edge in range(len(segment["edges"]) - 1):
            left, right = segment["edges"][edge], segment["edges"][edge + 1]
            inside = [word for word in words
                      if top - 1.0 <= (word["top"] + word["bottom"]) / 2 <= bottom + 1.0
                      and left - 1.0 <= word["x0"] < right - 1.0]
            cells.append(flat("".join(word["text"] for word in sorted(inside, key=lambda word: (word["x0"], word["top"])))))
        out.append(cells)
    return out


def all_bands(name: str) -> list:
    """整份 PDF 按先后读出来的带：(页号, 带号, 这一带的格子)。"""
    out = []
    for page_index, page in enumerate(pages_of(name)):
        for segment in segments(page):
            for row_index, row in enumerate(bands(page, segment)):
                out.append((page_index, row_index, row))
    return out


def expected(name: str) -> list:
    """排片单里按先后该出现的行（带表名的标题行不参与比对，只看表）。"""
    out = []
    for item in parse_plan(name):
        if item[0] == "表":
            out += [row for row in item[3]]
    return out


def align(name: str) -> list:
    """纸上的带 vs 文件里的行：只允许"跨页那一页开头多出一遍表头"。"""
    header_of = {}
    for item in parse_plan(name):
        if item[0] == "表" and item[2] and item[3]:
            header_of[tuple(item[3][0])] = True
    want = expected(name)
    got = [row for _page, _index, row in all_bands(name)]
    problems = []
    walk = 0
    for index, row in enumerate(got):
        if walk < len(want) and row == want[walk]:
            walk += 1
            continue
        if tuple(row) in header_of:            # 翻页后重复的表头，允许
            continue
        problems.append("纸上第 %d 带 %s 在文件里这个位置不是它（该是 %s）" % (index + 1, row, want[walk:walk + 1]))
        walk += 1
    if walk != len(want):
        problems.append("文件里还有 %d 行没在纸上读到（纸上 %d 带）" % (len(want) - walk, len(got)))
    return problems


def main() -> int:
    missing = [name for name in ALL_PRODUCTS
               if not os.path.isfile(os.path.join(BUILD, "%s.pdf" % name))
               or os.path.getsize(os.path.join(BUILD, "%s.pdf" % name)) == 0
               or not os.path.isfile(os.path.join(BUILD, "%s.notes.txt" % name))]
    if missing:
        print("红的是产物不在：%s" % "、".join(missing))
        print("先跑 tools/make_pdf_table_fixtures.py、./gradlew :core:jar，再用 tools/pdfprobe/run.sh "
              "pdftable（排片单那份）与 xlsxpdf（应用那条链子那份）出 PDF")
        return 1

    docs = {name: pages_of(name) for name in ALL_PRODUCTS}
    check(1, "六份 PDF（每样两份）与说明都在且非空；60 行那份真的翻了页",
          len(docs["long"]) >= 2 and all(docs[name] for name in FIXTURES),
          "页数 book=%d wide=%d long=%d" % (len(docs["book"]), len(docs["wide"]), len(docs["long"])))

    order = []
    cells = []
    for name in FIXTURES:
        order += ["%s：%s" % (name, item) for item in align(name)]
        for item in parse_plan(name):
            if item[0] != "表":
                continue
            _, title, _header, rows = item
            # 逐格：文件里的每一行，纸上必须有一带的格子顺序与内容完全一样
            paper_rows = {tuple(row) for _page, _index, row in all_bands(name)}
            for row_index, row in enumerate(rows):
                if tuple(row) not in paper_rows:
                    cells.append("%s/%s 第 %d 行没整行读到：%s" % (name, title, row_index + 1, row))
    check(2, "纸上按先后读出来的每一带与文件里的行序对得上（只许翻页重复表头）", not order, "; ".join(order[:2]))
    check(3, "逐格核位置：文件里的每一行都在纸上同一带、同一列顺序里", not cells, "; ".join(cells[:2]))

    grid = []
    heavy = []
    for name in FIXTURES:
        tables = [item for item in parse_plan(name) if item[0] == "表"]
        for page_index, page in enumerate(docs[name]):
            horizontals = [line for line in page.lines if abs(line["top"] - line["bottom"]) < 0.6]
            for segment in segments(page):
                edges = segment["edges"]
                read = bands(page, segment)
                columns = len(edges) - 1
                # 这一带的每一行都要是某张表里的行，而那张表的列数就是纸上读出来的列数。
                # 少一条竖线 → 最右那一列没了区间 → 这一带读出来不是文件里那一行
                owners = [item for item in tables
                          if max(len(row) for row in item[3]) == columns
                          and set(map(tuple, item[3])) >= {tuple(band) for band in read}]
                if not owners and any(any(cell for cell in band) for band in read):
                    grid.append("%s 第%d页那一带（%d 列）不属于排片单里的任何一张表：%s"
                                % (name, page_index + 1, columns, read[:1]))
                # 框要闭合：顶上一条、底下一条，每条横线都贯穿这段表
                inside = [line for line in horizontals if segment["top"] - 1.0 <= line["top"] <= segment["bottom"] + 1.0]
                if not any(abs(line["top"] - segment["top"]) <= 1.0 for line in inside):
                    grid.append("%s 第%d页这段表顶上的线没画" % (name, page_index + 1))
                if not any(abs(line["top"] - segment["bottom"]) <= 1.0 for line in inside):
                    grid.append("%s 第%d页这段表底下的线没画" % (name, page_index + 1))
                short = [line for line in inside if line["x0"] > edges[0] + 1.0 or line["x1"] < edges[-1] - 1.0]
                if short:
                    grid.append("%s 第%d页有横线只画到 x=%.1f，没贯穿到 %.1f"
                                % (name, page_index + 1, short[0]["x1"], edges[-1]))
            if not any(segments(page) for page in docs[name]):
                grid.append("%s 纸上一条竖线都没有 —— 一张表的框都没画出来" % name)
    for name in FIXTURES:
        for item in parse_plan(name):
            if item[0] != "表":
                continue
            _, title, header, rows = item
            mine = [segment for page in docs[name] for segment in segments(page)
                    if len(segment["edges"]) - 1 == max(len(row) for row in rows)]
            counted = sum(len(segment["heavy"]) for segment in mine)
            if header and counted == 0:
                heavy.append("%s/%s 表头那条重线没画出来" % (name, title))
            if not header and counted:
                heavy.append("%s/%s 没写表头却画了 %d 条重线" % (name, title, counted))
    check(4, "网格闭合且归得回：每一带都属于排片单里的某张表、列数一致，框上下有线、横线贯穿", not grid, "; ".join(grid[:2]))
    check(5, "表头下面那条线比别的重；没写表头的表不许有重线", not heavy, "; ".join(heavy[:2]))

    # 只有一页时不硬取第二页：那要让"没翻页"归到第 1 条，这条自己报"没翻过去"
    second = docs["long"][1] if len(docs["long"]) > 1 else None
    first = segments(second)[0] if second is not None and segments(second) else None
    head_row = bands(second, first)[0] if first else []
    check(6, "跨页那份：第二页第一带就是表头那一行", head_row == ["表头甲", "表头乙", "表头丙"],
          "第二页第一带=%s（long 共 %d 页）" % (head_row, len(docs["long"])))

    with open(os.path.join(BUILD, "book.notes.txt"), encoding="utf-8") as handle:
        notes = handle.read()
    pressed = re.findall(r"(?:「(.+?)」)?第 (\d+) 列从 \d+pt 压到 (\d+)pt", notes)
    honest = []
    columns_of = {}
    for item in parse_plan("book"):
        if item[0] == "表":
            columns_of[item[1]] = max(len(row) for row in item[3])
    for title, column, to_pt in pressed:
        index = int(column) - 1
        want_columns = columns_of.get(title) if title else None
        target = None
        for page in docs["book"]:
            for segment in segments(page):
                columns = len(segment["edges"]) - 1
                if len(segment["edges"]) > index + 1 and (want_columns is None or columns == want_columns):
                    target = segment
                    break
            if target:
                break
        if target is None:
            honest.append("notes 点了%s第 %s 列，纸上没这张表/这么多列" % ("「%s」" % title if title else "", column))
            continue
        measured = target["edges"][index + 1] - target["edges"][index]
        if abs(measured - float(to_pt)) > 3.0:
            honest.append("notes 说%s第 %s 列压到 %spt，量出来 %.1fpt"
                          % ("「%s」" % title if title else "", column, to_pt, measured))
    check(7, "压窄的说明诚实：notes 说压到多少，纸上量出来就是多少", bool(pressed) and not honest,
          "; ".join(honest) or "notes=%s" % notes.strip().replace("\n", " / ")[:90])

    # parse_plan 把格子抹平了（比内容要用），这一条直接读原始排片单里那句带硬换行的格子
    with open(os.path.join(BUILD, "book.plan.tsv"), encoding="utf-8") as handle:
        raw_cell = next(line.split("\t")[1] for line in handle.read().splitlines()
                        if line.startswith("3\t") and "\t" in line)
    pieces = [flat(piece) for piece in raw_cell.split("\\n") if piece.strip()]
    words = [word for page in docs["book"] for word in words_on(page)]
    needles = [piece[:6] for piece in pieces if piece]
    both = all(any(needle in word["text"] for word in words) for needle in needles)
    tops = set()
    for word in words:
        if any(needle in word["text"] for needle in needles):
            tops.add(round(word["top"], 1))
    lines_apart = len(tops) >= 2
    check(8, "格子里的硬换行写在两条基线上（不是挤成一行）", both and lines_apart,
          "片段=%s 两段都在=%s 基线=%s" % (needles, both, sorted(tops)))

    title_bad = []
    for name in FIXTURES:
        body = min((word["size"] for page in docs[name] for word in words_on(page)), default=0.0)
        tables = [item for item in parse_plan(name) if item[0] == "表"]
        for page_index, page in enumerate(docs[name]):
            words_here = words_on(page)
            on_page = segments(page)
            for item in tables:
                title, columns = item[1], max(len(row) for row in item[3])
                mine = [segment for segment in on_page if len(segment["edges"]) - 1 == columns]
                found = [word for word in words_here if flat(title) in flat(word["text"])]
                if not found or not mine:
                    continue
                if not any(word["bottom"] <= segment["top"] for word in found for segment in mine):
                    title_bad.append("%s 第%d页「%s」没排在它自己那张表上面" % (name, page_index + 1, title))
                if max(word["size"] for word in found) <= body + 0.5:
                    title_bad.append("%s/%s 表名字号 %.1f 不比正文大" % (name, title, max(word["size"] for word in found)))
    check(9, "表名那行是表名、排在它自己那张表上面，且字号比正文大", not title_bad, "; ".join(title_bad[:2]))

    flat_all = flat("".join(word["text"] for page in docs["book"] for word in words_on(page)))
    empties = [item[1] for item in parse_plan("book") if item[0] == "空"]
    sentence = next((item[1] for item in parse_plan("book") if item[0] == "文"), "")
    last_bottom = max((segment["bottom"] for page in docs["book"] for segment in segments(page)), default=0.0)
    sentence_words = [word for page in docs["book"] for word in words_on(page) if sentence[:5] in word["text"]]
    after_tables = bool(sentence_words) and max(word["top"] for word in sentence_words) > last_bottom
    check(10, "空表不占一节（标题不在纸上），表外那句话排在所有表后面",
          bool(empties) and all(flat(title) not in flat_all for title in empties) and after_tables,
          "空表标题在纸上=%s 句子在表后=%s（最后一张表底 %.1f）"
          % ([title for title in empties if flat(title) in flat_all], after_tables, last_bottom))

    workbook = openpyxl.load_workbook(os.path.join(BUILD, "book.xlsx"))
    cells_only = sum(len(flat("" if cell is None else str(cell)))
                     for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True) for cell in row)
    # 纸上的字不止格子里那些：每张表上面那行表名、以及夹在表中间的一句话都要印，比对要按整份算
    plan_chars = 0
    for item in parse_plan("book"):
        if item[0] == "表":
            plan_chars += len(flat(item[1])) + sum(len(cell) for row in item[3] for cell in row)
        elif item[0] == "文":
            plan_chars += len(flat(item[1]))
    import fitz                                          # MuPDF：与 pdfminer 不是同一套解析
    document = fitz.open(os.path.join(BUILD, "book.pdf"))
    got = len(flat("".join(page.get_text() for page in document)))
    document.close()
    # poppler 那一系要显式要 UTF-8：这台机器上 `pdftotext f.pdf -` 默认按 Latin-1 出，
    # 中文全被丢掉（0 个字），差点被我当成"产物缺中文文字层"。加 -enc UTF-8 后与 MuPDF 一字不差。
    pop = subprocess.run(["pdftotext", "-enc", "UTF-8", os.path.join(BUILD, "book.pdf"), "-"],
                         capture_output=True)
    pop_chars = len(flat(pop.stdout.decode("utf-8", "replace")))
    check(11, "MuPDF 与 poppler（两套独立解析）读回来的字数都与原件相当，且两家用词数一致",
          cells_only < plan_chars and plan_chars * 0.97 <= got <= plan_chars * 1.03 and
          plan_chars * 0.97 <= pop_chars <= plan_chars * 1.03 and abs(got - pop_chars) <= plan_chars * 0.02,
          "应有 %d 字（格子 %d），MuPDF %d、poppler %d" % (plan_chars, cells_only, got, pop_chars))

    outside = []
    narrow = []
    for name in ALL_PRODUCTS:
        for page_index, page in enumerate(docs[name]):
            for word in words_on(page):
                if word["x1"] > PAGE_WIDTH - MARGIN + 2.0 or word["x0"] < MARGIN - 2.0:
                    outside.append("%s 第%d页「%s」x=%.1f..%.1f" % (name, page_index + 1, word["text"], word["x0"], word["x1"]))
            for segment in segments(page):
                for index in range(len(segment["edges"]) - 1):
                    if segment["edges"][index + 1] - segment["edges"][index] < 6.0:
                        narrow.append("%s 第 %d 列只有 %.1fpt" % (name, index + 1, segment["edges"][index + 1] - segment["edges"][index]))
    check(12, "没有字越出页边距；压窄也不许压成零宽", not outside and not narrow,
          "; ".join((outside + narrow)[:2]))

    blank = ["%s 第%d页" % (name, index + 1) for name in ALL_PRODUCTS
             for index, page in enumerate(docs[name]) if not words_on(page)]
    check(13, "没有空白页（排了几页就几页有东西）", not blank, "; ".join(blank[:2]))

    # 14/15 走的是**应用那条链子**：真 .xlsx 由 :core 的 XlsxSheets.book 读、parts 成树、PrintPlanner 排页，
    # 只有取部件字节那一步换成 JDK 的 ZipFile。排片单（openpyxl 读的）是这一条的参照物。
    app_bad = []
    for name in FIXTURES:
        want = expected(name)
        got = [row for _page, _index, row in all_bands(name + "-app")]
        if got != want:
            at = next((index for index in range(max(len(got), len(want)))
                       if got[index:index + 1] != want[index:index + 1]), min(len(got), len(want)))
            app_bad.append("%s-app 第 %d 带就不一样：纸上 %s / 文件里 %s（共 纸上%d 带、文件%d 行）"
                           % (name, at + 1, got[at:at + 1], want[at:at + 1], len(got), len(want)))
    check(14, "从真 .xlsx 走应用那条链子：纸上的带与 openpyxl 那份排片单逐带一致（空表不占节、也不重复表头）",
          not app_bad, "; ".join(app_bad[:2]))

    say_bad = []
    for name in FIXTURES:
        workbook = openpyxl.load_workbook(os.path.join(BUILD, "%s.xlsx" % name))
        filled = []
        for sheet in workbook.worksheets:
            rows = [row for row in sheet.iter_rows(values_only=True) if any(c is not None for c in row)]
            if rows:
                filled.append((sheet.title, len(rows), max(len(row) for row in rows)))
        with open(os.path.join(BUILD, "%s-app.notes.txt" % name), encoding="utf-8") as handle:
            say = handle.read()
        want = "%d 张表 · %d 行" % (len(filled), sum(count for _t, count, _c in filled))
        if want not in say:
            say_bad.append("%s-app 的说明里没有「%s」：%s" % (name, want, say.strip().replace("\n", " / ")[:80]))
        for title, count, columns in filled:
            if "%s：%d 行 × %d 列" % (title, count, columns) not in say:
                say_bad.append("%s-app 少了「%s：%d 行 × %d 列」" % (name, title, count, columns))
        for title in [s.title for s in workbook.worksheets
                      if not any(any(c is not None for c in row) for row in s.iter_rows(values_only=True))]:
            # 空表的名字**该**出现在"几张表是空的、没占一节"那句里（那是交代），
            # 但不该被当成一张表进行列清单（"空的：N 行 × M 列"）
            if re.search(re.escape(title) + r"：\d+ 行", say):
                say_bad.append("%s-app 把空表「%s」也列进行列清单了" % (name, title))
    check(15, "应用那条链子的说明诚实：notes 里几张表、每张几行几列，与 openpyxl 数出来的一致",
          not say_bad, "; ".join(say_bad[:2]))

    failed = [item for item in results if not item[2]]
    print("\n共 %d 条，%d 条红" % (len(results), len(failed)))
    for number, name, _ in failed:
        print("  红的是第 %d 条：%s" % (number, name))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
