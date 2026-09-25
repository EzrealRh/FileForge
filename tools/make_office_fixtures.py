#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""造 OOXML 夹具：手工拼出 Word / PowerPoint 的最小合法包。

本机没装 Word 也没装 python-docx，所以部件 XML 全部照 ECMA-376 的树形手写，
再用标准库 zipfile 装成包 —— 这样"包能被 zipfile 打开"本身就是一条外部约束。

三份夹具各自要卡住的东西不同：
  prose.docx   只有段落、修订、域、超链接 —— 这一份要能和 pandoc 抽出的文字**逐段相同**
  tables.docx  表格、文本框、图片、脚注引用 —— 这些 pandoc 有它自己的规矩，只跟自家期望比
  deck.pptx    两页幻灯片，每页多个形状

另外还压一份 book.xlsx 交给 **openpyxl**（真实现压出来的比手写的更像会碰到的文件），
参照值 `book.xlsx.truth` 记的是 openpyxl **读回来**看到的格子，不是我以为写进去的。

输出到 core/src/test/resources/office/。用法：python tools/make_office_fixtures.py
"""
import datetime
import io
import os
import zipfile
from xml.etree import ElementTree

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..", "core", "src", "test", "resources", "office"
)

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R_NS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
WP_NS = 'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
PIC_NS = 'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'

XML_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def run(text, preserve=False):
    """一个文本 run。preserve 决定边上的空格保不保 —— Word 需要时才写这个属性。"""
    space = ' xml:space="preserve"' if preserve else ""
    return "<w:r><w:t%s>%s</w:t></w:r>" % (space, text)


def para(*parts):
    return "<w:p>%s</w:p>" % "".join(parts)


def document(body):
    return XML_DECL + ('<w:document %s %s %s %s %s><w:body>%s<w:sectPr/></w:body></w:document>'
                       % (W_NS, R_NS, WP_NS, A_NS, PIC_NS, body))


CONTENT_TYPES_DOCUMENT = XML_DECL + (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    '<Override PartName="/word/footnotes.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
    "</Types>"
)

ROOT_RELS = XML_DECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/></Relationships>'
)

DOC_RELS = XML_DECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId10" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
    'Target="https://example.com/" TargetMode="External"/>'
    "</Relationships>"
)

FOOTNOTES = XML_DECL + (
    '<w:footnotes %s><w:footnote w:type="separator" w:id="-1"/><w:footnote w:type="continuationSeparator" '
    'w:id="0"/><w:footnote w:id="1"><w:p>%s</w:p></w:footnote></w:footnotes>'
    % (W_NS, run("脚注里的一句话"))
)


def prose_document():
    """段落、修订、域、超链接。每段都单独可核对。"""
    body = "".join([
        para(run("项目清单")),
        para(run("第一批：", True), run("甲"), "<w:tab/>", run("乙")),
        para(run("备注 A&#8212;B &amp; &lt;tag&gt; &#8220;引号&#8221;")),
        # 修订：ins 是新增的字，del 里的 delText 是被删掉的字（不该出现在结果里）
        para('<w:ins w:id="1" w:author="谁"><w:r><w:t>新增保留</w:t></w:r></w:ins>'
             '<w:del w:id="2" w:author="谁"><w:r><w:delText>删掉不该出现</w:delText></w:r></w:del>'),
        # 域：instrText 是代码，separate 之后的 run 才是显示结果
        para('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
             '<w:r><w:instrText> PAGE </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + run("7") +
             '<w:r><w:fldChar w:fldCharType="end"/></w:r>'),
        # 超链接在内联层：它只能是 w:p 的孩子，跟 w:r 平级（挂在 body 上就是不合法的，pandoc 会直接丢掉）
        '<w:p><w:hyperlink r:id="rId10" w:history="1">%s</w:hyperlink></w:p>' % run("官网"),
        "<w:p/>",
        para(run("结尾")),
    ])
    return document(body)


PROSE_EXPECT = [
    "项目清单",
    "第一批：甲\t乙",
    "备注 A—B & <tag> “引号”",
    "新增保留",
    "7",
    "官网",
    "",                                        # 夹具里那个空段落
    "结尾",
]


def tables_document():
    """表格、文本框、图片、脚注引用 —— 这些各家规矩不同，只跟自家期望比。"""
    table = (
        "<w:tbl><w:tblPr/><w:tr><w:tc><w:tcPr/>%s</w:tc><w:tc><w:tcPr/>%s</w:tc>"
        "<w:tc><w:tcPr/>%s</w:tc></w:tr>"
        "<w:tr><w:tc><w:tcPr/>%s</w:tc><w:tc><w:tcPr/>%s</w:tc><w:tc><w:tcPr/><w:p/></w:tc></w:tr></w:tbl>"
        % (
            para(run("名称")), para(run("数量")), para(run("备注")),
            para(run("苹果")), para(run("3")),
        )
    )
    cell_two_paras = (
        "<w:tbl><w:tr><w:tc><w:tcPr/>%s%s</w:tc><w:tc><w:tcPr/>%s</w:tc></w:tr></w:tbl>"
        % (para(run("格内第一段")), para(run("格内第二段")), para(run("右格")))
    )
    drawing = (
        '<w:r><w:drawing><wp:inline><wp:extent cx="1" cy="1"/><a:graphic><a:graphicData>'
        '<pic:pic><pic:blipFill><a:blip r:embed="rId5"/></pic:blipFill></pic:pic>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
    )
    text_box = (
        '<w:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice xmlns:v="urn:schemas-microsoft-com:vml" Requires="v">'
        '<w:drawing><wp:inline><w:txbxContent>%s</w:txbxContent></wp:inline></w:drawing>'
        "</mc:Choice></mc:AlternateContent></w:r>" % para(run("文本框里的字"))
    )
    footnote_ref = '<w:r><w:footnoteReference w:id="1"/></w:r>'
    body = "".join([
        para(run("前言")),
        table,
        cell_two_paras,
        para(run("带图的"), drawing),
        para(run("带脚注的"), footnote_ref, run("这句话")),
        para(run("框外"), text_box),
        "<w:p><w:pPr><w:sectPr><w:headerReference r:id=\"rId7\" w:type=\"default\"/>"
        "<w:footerReference r:id=\"rId8\" w:type=\"default\"/></w:sectPr></w:pPr></w:p>",
        para(run("结尾")),
    ])
    return document(body)


TABLES_EXPECT = [
    "前言",
    "名称\t数量\t备注",
    "苹果\t3\t",
    "格内第一段 格内第二段\t右格",
    "带图的",
    "带脚注的这句话",
    "框外文本框里的字",
    "",                                        # 带页眉页脚引用的那段是空的
    "结尾",
]


def slide(*shapes):
    return (XML_DECL +
            '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            "<p:cSld><p:spTree>%s</p:spTree></p:cSld></p:sld>" % "".join(shapes))


def shape(*paragraphs):
    return "<p:sp><p:nvSpPr/><p:spPr/><p:txBody>%s</p:txBody></p:sp>" % "".join(paragraphs)


def a_para(text):
    return "<a:p><a:r><a:t>%s</a:t></a:r></a:p>" % text


SLIDE_ONE = slide(shape(a_para("第一页标题"), a_para("副标题")), shape(a_para("右下角备注")))
SLIDE_TWO = slide(
    "<p:graphicFrame><a:graphic><a:graphicData><a:tbl>"
    "<a:tr><a:tc><a:txBody>%s</a:txBody></a:tc><a:tc><a:txBody>%s</a:txBody></a:tc></a:tr>"
    "<a:tr><a:tc><a:txBody>%s</a:txBody></a:tc><a:tc><a:txBody>%s</a:txBody></a:tc></a:tr>"
    "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
    % (a_para("列一"), a_para("列二"), a_para("1"), a_para("2")),
    # 形状必须在 p:sld 这棵根里面：写成 slide(...) + shape(...) 会变成两份根，pandoc 直接报"根之后还有内容"
    shape('<a:p><a:fld id="{B}" type="slidenumber"><a:t>2</a:t></a:fld></a:p>'),
)

DECK_EXPECT = [
    "第一页标题",
    "副标题",
    "右下角备注",
    "列一\t列二",
    "1\t2",
    "2",
]

CONTENT_TYPES_DECK = XML_DECL + (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/ppt/presentation.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
    '<Override PartName="/ppt/slides/slide1.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
    '<Override PartName="/ppt/slides/slide2.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
    "</Types>"
)

# 根关系指向哪份主部件是包级别的规矩：演示文稿指 ppt/presentation.xml。
# （这条早先抄了文档包的那份，pandoc 立刻按关系去找 word/document.xml 然后报"部件不存在"。）
DECK_ROOT_RELS = XML_DECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="ppt/presentation.xml"/></Relationships>'
)

DECK_RELS = XML_DECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId2" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" '
    'Target="slides/slide1.xml"/>'
    '<Relationship Id="rId3" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" '
    'Target="slides/slide2.xml"/></Relationships>'
)

PRESENTATION = XML_DECL + (
    '<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<p:sldIdLst><p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId3"/></p:sldIdLst>'
    "</p:presentation>"
)


def write_zip(name, parts):
    path = os.path.join(OUT, name)
    # 先体检再生成：每个部件都得能被 ElementTree 解析。夹具自己坏了，测出来的"一致"毫无意义
    for part, payload in parts:
        ElementTree.fromstring(payload.encode("utf-8"))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for part, payload in parts:
            package.writestr(part, payload)
    # 自己先开一遍：能被 zipfile 打开才算一份包
    with zipfile.ZipFile(path) as check:
        assert check.testzip() is None, name
        names = check.namelist()
    print("  %s：%d 个部件，%d 字节" % (name, len(names), os.path.getsize(path)))
    return names


def build_book(path):
    """用 openpyxl 压一份真 xlsx 当参照物：表名、日期、自定义日期格式、稀疏行、空表都有。

    故意不放公式：openpyxl 压出来的文件里没有"算过的结果"，那一格的取值规矩两端本来就不同，
    由 Kotlin 侧自己用手写部件的单测钉住。
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        print("  跳过 book.xlsx：本机没有 openpyxl（pip install openpyxl 后重跑本脚本）")
        return False
    wb = Workbook()
    bill = wb.active
    bill.title = "费用"
    bill["A1"] = "项目"
    bill["B1"] = "金额"
    bill["C1"] = "发生日"
    bill["A2"] = "打车"
    bill["B2"] = 38.5
    bill["C2"] = datetime.date(2023, 5, 1)
    bill["C3"] = datetime.datetime(2024, 2, 29, 8, 30, 15)
    bill["C3"].number_format = "yyyy-mm-dd hh:mm:ss"
    bill["A4"] = "含中文的格子，带逗号,"
    bill["C4"] = datetime.date(2021, 3, 4)
    bill["C4"].number_format = 'yyyy"年"m"月"d"日"'      # 带字面文字的自定义格式：日期码在引号外面
    bill["D5"] = "第五行只有 D 列有东西"                     # 行号与列号都跳过前面
    bill["E6"] = 0.25
    bill["B7"] = "纯文字"
    bill["C7"].number_format = "0.00"                       # 没赋值的格子：样式是数字，值还是没有
    empty = wb.create_sheet("空表")                          # 一张完全空的表
    hidden = wb.create_sheet("第三张")
    hidden["A1"] = "顺序要按工作簿声明的来"
    wb.save(path)
    print("  book.xlsx：由 openpyxl 压出，%d 字节" % os.path.getsize(path))
    return True


def book_truth(path):
    """参照值 = openpyxl **读回来**看到的东西，按"一行一记录"写；日期一律展成 ISO 写法。"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    out = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([render_cell(v) for v in row])
        while rows and all(not cell for cell in rows[-1]):
            rows.pop()
        cells = ["# " + sheet.title]
        cells += ["\t".join(row) for row in rows]
        out.append("\n".join(cells))
    return out


def render_cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        # openpyxl 把日期一律读成 datetime（Python 的类型系统决定的）：零点就还原成"只有日期"，
        # 那才是文件里存的样子，也是 Excel 显示的样子
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M:%S")
    return str(value)


def main():
    os.makedirs(OUT, exist_ok=True)
    write_zip("prose.docx", [
        ("[Content_Types].xml", CONTENT_TYPES_DOCUMENT),
        ("_rels/.rels", ROOT_RELS),
        ("word/_rels/document.xml.rels", DOC_RELS),
        ("word/document.xml", prose_document()),
        ("word/footnotes.xml", FOOTNOTES),
    ])
    write_zip("tables.docx", [
        ("[Content_Types].xml", CONTENT_TYPES_DOCUMENT),
        ("_rels/.rels", ROOT_RELS),
        ("word/_rels/document.xml.rels", DOC_RELS),
        ("word/document.xml", tables_document()),
        ("word/footnotes.xml", FOOTNOTES),
    ])
    write_zip("deck.pptx", [
        ("[Content_Types].xml", CONTENT_TYPES_DECK),
        ("_rels/.rels", DECK_ROOT_RELS),
        ("ppt/_rels/presentation.xml.rels", DECK_RELS),
        ("ppt/presentation.xml", PRESENTATION),
        ("ppt/slides/slide1.xml", SLIDE_ONE),
        ("ppt/slides/slide2.xml", SLIDE_TWO),
    ])
    for name, lines in (
        ("prose.docx", PROSE_EXPECT), ("tables.docx", TABLES_EXPECT), ("deck.pptx", DECK_EXPECT)
    ):
        # 期望值按"一行一段"写成 .expect：两端的比对单位就是行，不各写一份解析器
        with io.open(os.path.join(OUT, name + ".expect"), "w", encoding="utf-8", newline="\n") as handle:
            handle.write("\n".join(lines) + "\n")
    book = os.path.join(OUT, "book.xlsx")
    if build_book(book):
        blocks = book_truth(book)
        with io.open(os.path.join(OUT, "book.xlsx.truth"), "w", encoding="utf-8", newline="\n") as handle:
            handle.write("\n".join(blocks) + "\n")
    print("写到", os.path.normpath(OUT))


if __name__ == "__main__":
    main()
