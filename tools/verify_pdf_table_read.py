"""PDF 表格读回的判据：纸上的行列能不能一字不差地还原成原件那张表。

四边都不沾的一条链子：

    openpyxl 读原件（真值） →  应用那条排版路印成 PDF（:core 的 PrintPlanner + PDFBox 画）
        →  应用那条读路读回格子（PDFBox 量词 + :core 的 PdfTableDetect / PdfTableSheets）
        →  与真值逐格比

第三方的独立读路（pdfplumber 的 extract_words）只核一件事：**我们选的每条列边界，
都得在它独立切出的词里找得到多行重复的词起点** —— 它在中文上会把一个词按字劈开（实测
"列10"读成"列"、"第 38 行"读成三格），所以拿它的表当等号反而会逼我们学它那份坏行为，
这一条只核"没凭空造列"。

跑法（工程根目录）：
    GRADLE_USER_HOME=D:/gradle-home ./gradlew :core:jar
    python tools/verify_pdf_table_read.py
"""

import os
import subprocess
import sys

import openpyxl
import pdfplumber

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BUILD = os.path.join("core", "build", "pdftable")
WORK = os.path.join("core", "build", "pdftable-read")
FONT = os.environ.get("PDF_TABLE_FONT", "C:/Windows/Fonts/simhei.ttf")
RUN = os.path.join("tools", "pdfprobe", "run.sh")

RESULTS = []


def report(name, ok, detail=""):
    RESULTS.append((name, ok))
    print("%s %s%s" % ("PASS" if ok else "FAIL", name, "" if not detail else "\n       " + detail))


def probe(*args):
    env = dict(os.environ, PDFPROBE_OUT=os.environ.get("PDFPROBE_OUT", "D:/fflogs/pdfprobe"))
    out = subprocess.run(["bash", RUN] + list(args), capture_output=True, text=True,
                         encoding="utf-8", errors="replace", env=env)
    if out.returncode != 0:
        raise RuntimeError("探针跑不动：%s\n%s" % (" ".join(args), (out.stderr or out.stdout)[-800:]))
    return out.stdout


def plan_truth(name):
    """从 .xlsx 用 openpyxl 现读真值：每张非空表 = 行 × 格的字符串矩阵。"""
    book = openpyxl.load_workbook(os.path.join(BUILD, name + ".xlsx"))
    sheets = []
    for title in book.sheetnames:
        rows = []
        for row in book[title].iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells or rows:
                rows.append(cells)
        if any(any(cell for cell in row) for row in rows):
            width = max(len(row) for row in rows)
            sheets.append([row + [""] * (width - len(row)) for row in rows])
    return sheets


def read_back(pdf, tag):
    """跑应用那条读路，把 sheets.tsv 收回来：[(表名, [格, …]), …]。"""
    out_dir = os.path.join(WORK, tag)
    probe("pdftables", pdf, out_dir)
    sheets = []
    with open(os.path.join(out_dir, "sheets.tsv"), encoding="utf-8") as handle:
        for line in handle:
            parts = line.rstrip("\n").split("\t")
            if len(parts) < 2:
                continue
            if not sheets or sheets[-1][0] != parts[0]:
                sheets.append((parts[0], []))
            sheets[-1][1].append([cell.replace("\\n", "\n") for cell in parts[1:]])
    return sheets


def flat(cells):
    """比内容时用：换行丢掉（同一格里软换行读回来会多一个 \\n，硬换行原件本来就有）。"""
    return "".join(cell.replace("\n", "") for cell in cells)


def normalize(rows):
    return [[cell.replace("\n", "").strip() for cell in row] for row in rows]


def compare(name, truth_sheets, got):
    detail = []
    if len(got) != len(truth_sheets):
        detail.append("表张数不同：真值 %d 张，读回 %d 张" % (len(truth_sheets), len(got)))
    for index, truth in enumerate(truth_sheets):
        if index >= len(got):
            detail.append("第 %d 张没读回来" % (index + 1))
            continue
        rows = normalize(got[index][1])
        want = normalize(truth)
        if len(rows) != len(want):
            detail.append("第 %d 张行数不同：真值 %d 行，读回 %d 行" % (index + 1, len(want), len(rows)))
        for row_index in range(min(len(rows), len(want))):
            if rows[row_index] != want[row_index]:
                detail.append("第 %d 张第 %d 行不等：\n         真 %r\n         读 %r"
                              % (index + 1, row_index + 1, want[row_index], rows[row_index]))
        if not detail and flat(sum(want, [])) != flat(sum(rows, [])):
            detail.append("第 %d 张的字对不上（行数格子数一样但内容有出入）" % (index + 1))
    report(name, not detail, "；".join(detail[:4]))


def cuts_of(pdf, tag):
    """同一批词的列边界：page, top, bottom, 行数, [边界…]。"""
    out_dir = os.path.join(WORK, tag)
    probe("pdftables", pdf, out_dir)
    rows = []
    path = os.path.join(out_dir, "cuts.tsv")
    if not os.path.exists(path):
        return rows
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            parts = line.rstrip("\n").split("\t")
            if len(parts) < 5:
                continue
            rows.append((int(parts[0]), float(parts[1]), float(parts[2]), int(parts[3]),
                         [float(x) for x in parts[4].split(",") if x]))
    return rows


def compare_columns(name, truth_sheets, got, want_columns):
    """窄到装不下的宽表（一格只够放两三个字，纸上是逐字往下堆的）不可能按行还原回去：
    独立读路 pdfplumber 在同样的纸上给出 27 行，我们给 25 行，真值是 3 行。
    所以对这一份只核"列找对没、字有没有整列搬齐"，行数不装成能核。"""
    detail = []
    if len(got) != 1:
        detail.append("该读回 1 张表，实际 %d 张" % len(got))
        report(name, False, "；".join(detail))
        return
    rows = [[cell.replace("\n", "").strip() for cell in row] for row in got[0][1]]
    if any(len(row) != want_columns for row in rows):
        detail.append("有的行不是 %d 格：%s" % (want_columns, sorted({len(row) for row in rows})))
    truth = [row for sheet in truth_sheets for row in sheet]
    for column in range(want_columns):
        mine = "".join(row[column] for row in rows if len(row) > column)
        want = "".join(row[column] if column < len(row) else "" for row in truth)
        # 真值那一列的字连起来，与读回来那一列的字连起来，必须一字不差（顺序也一致）
        flat_want = "".join(want.split())
        flat_mine = "".join(mine.split())
        if flat_want != flat_mine:
            detail.append("第 %d 列内容不齐：真 %r 读 %r" % (column + 1, flat_want[:40], flat_mine[:40]))
    report(name, not detail, "；".join(detail[:3]))
    if not detail:
        print("       （这一份只核列不核行：纸上每格只剩两三个字的宽度，行数没法还原，独立读路同样还原不了）")


def repeated_starts(pdf, min_lines=3):
    """pdfplumber 独立切的词里，被至少 min_lines 行共用（±2pt）的词起点。"""
    starts = []
    with pdfplumber.open(pdf) as book:
        for page in book.pages:
            per_line = {}
            for word in page.extract_words(use_text_flow=False, keep_blank_chars=False):
                per_line.setdefault(round(word["top"], 1), set()).add(round(word["x0"]))
            for seen in per_line.values():
                starts.extend(seen)
    clusters = []
    for at in sorted(starts):
        if clusters and abs(clusters[-1][-1] - at) <= 2:
            clusters[-1].append(at)
        else:
            clusters.append([at])
    return [sum(group) / len(group) for group in clusters if len(group) >= min_lines]


def main():
    os.makedirs(WORK, exist_ok=True)
    for name in ("book", "wide", "long"):
        probe("xlsxpdf", FONT, os.path.join(BUILD, name + ".xlsx"), os.path.join(BUILD, name + "-app.pdf"))
        probe("pdftable", FONT, os.path.join(BUILD, name + ".plan.tsv"), os.path.join(BUILD, name + ".pdf"))

    # 1. 三份原件：应用印出来的 PDF 再让应用读回来，必须与 openpyxl 读的原件对得上
    #    （book / long 逐格逐行；wide 只核列 —— 理由见 compare_columns）
    for name in ("book", "long"):
        compare("%s：印一遍再读回来，与 openpyxl 读的原件逐格相同" % name,
                plan_truth(name), read_back(os.path.join(BUILD, name + "-app.pdf"), name + "-app"))
    compare_columns("wide：26 列一列不多一列不少，每列的字按顺序整列搬齐",
                    plan_truth("wide"), read_back(os.path.join(BUILD, "wide-app.pdf"), "wide-app"), 26)
    # 排片单那条也照原件核一遍：那份 PDF 每页顶上会重印一行表头，读回来必须只算一次
    compare("long：排片单那条印的那份读回来也与原件逐格相同（跨页重印的表头只算一次）",
            plan_truth("long"), read_back(os.path.join(BUILD, "long.pdf"), "long-plan"))

    # 2. 两条排版路（应用那条与排片单那条）读回来要给出同样的格子，否则说明有一边在挑数据
    for name in ("book", "long"):
        via_app = normalize(sum([r for _, r in read_back(os.path.join(BUILD, name + "-app.pdf"), name + "-app")], []))
        via_plan = normalize(sum([r for _, r in read_back(os.path.join(BUILD, name + ".pdf"), name + "-plan")], []))
        # 排片单那份多印了两张表之间的一句话，比内容时只比"每张表里的字连起来"
        report("%s：应用那条链子与排片单那条读回同样的格子" % name,
               [flat(r) for r in via_app if any(r)] == [flat(r) for r in via_plan if any(r)],
               "app=%d 行 plan=%d 行" % (len(via_app), len(via_plan)))

    # 3. 独立读路核"没凭空造列"：我们选的每条列边界，都得是 pdfplumber 自己切的词里
    #    被多行重复的那个起点（±2pt）；每一页最宽那张表的列数说死（同一页另有窄表是正常的）
    for name, want_columns in (("book-app", 4), ("wide-app", 26), ("long-app", 3)):
        pdf = os.path.join(BUILD, name + ".pdf")
        allowed = repeated_starts(pdf)
        detail = []
        by_page = {}
        for page, top, bottom, lines, boundaries in cuts_of(pdf, name + "-cuts"):
            for at in boundaries:
                if not any(abs(at - candidate) <= 2.0 for candidate in allowed):
                    detail.append("第 %d 页 %.1f 这条列边界在独立读路里找不到重复的词起点" % (page, at))
            by_page[page] = max(by_page.get(page, 0), len(boundaries) + 1)
        if not by_page:
            detail.append("没读出任何一段表")
        for page, columns in sorted(by_page.items()):
            if columns != want_columns:
                detail.append("第 %d 页最宽那张读成 %d 列，独立看这份纸该是 %d 列" % (page, columns, want_columns))
        report("%s：列边界都落在独立读路重复出现的词起点上（最宽 %d 列）" % (name, want_columns),
               not detail, "；".join(detail[:3]))

    # 4. 负例：一页只有整句正文，没有表 —— 不许硬凑出一张表，且要说清为什么
    prose = os.path.join(WORK, "prose.md")
    with open(prose, "w", encoding="utf-8") as handle:
        handle.write("\n\n".join([
            "第一段讲的是本机处理不上传这件事，句子写得长一些，好在排的时候自然折成好几种宽度。",
            "第二段接着说换行标点。收尾标点不许顶到行首，长串要能在任意处断开，这些都跟表格无关。",
            "第三段随便再写一点，凑出几行长短不齐的字，让任何想找列的人都找不到一处稳定的起点。",
        ]) + "\n")
    probe("xlsxpdf", FONT, prose, os.path.join(WORK, "prose.pdf"))
    out_dir = os.path.join(WORK, "prose-read")
    log = probe("pdftables", os.path.join(WORK, "prose.pdf"), out_dir)
    got = read_back(os.path.join(WORK, "prose.pdf"), "prose-read2")
    report("整页正文不产表", len(got) == 0, "读回 %d 张：%s" % (len(got), [name for name, _ in got]))
    report("不产表时把原因说出来", ("没当成表" in log) or ("对不齐" in log), log.strip()[-200:])

    red = [name for name, ok in RESULTS if not ok]
    print("\n%d 条判据，%d 条红" % (len(RESULTS), len(red)))
    for name in red:
        print("  RED:", name)
    sys.exit(1 if red else 0)


if __name__ == "__main__":
    main()
