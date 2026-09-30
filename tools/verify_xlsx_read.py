#!/usr/bin/env python3
"""
判 xlsx 的"一批表 → 文档树"那一步与它的三条产物（转网页 / 转 Markdown / 写成 Word）。

跑法（先跑 JVM 那边把产物落出来）：
    GRADLE_USER_HOME=D:/gradle-home ./gradlew :core:test --tests "com.fileforge.core.XlsxTest" --rerun
    python tools/verify_xlsx_read.py

参照物是 **openpyxl**：它按坐标自己把格子摆一遍（`<c r="C3">` 前面没有 A、B 也是 C3），
pandoc 不读 xlsx，所以这里 pandoc 只用来读**我们的产物**（Markdown / 网页 / Word 三份），
把"别人读我们写出去的东西"与 openpyxl 读原件的结果对齐。

十条判据：
  1 openpyxl 读得动这份 xlsx（读不动就是夹具坏了，红要指名）
  2 三条产物都在且非空
  3 一节一张表：产物里的二级标题序列 == openpyxl 眼里"有格子的表"，且顺序照工作簿声明
  4 整张格子逐格对：openpyxl 按坐标摆的网格与三份产物读回来的表一致（三条产物彼此也一致）
  5 网格尺寸：行数与列数 == openpyxl 的最大行号与最大列号（尾巴空行不被截、空列要补上）
  6 日期不是序列号：openpyxl 给 datetime 的那些格，产物里是 ISO 写法；有时间的要把时间带着
  7 数字与文字照原样：38.5 / 0.25 / 带逗号的那格在产物里还是那串字
  8 表头那一行留空：普通区域没声明首行是表头，产物里就不该有表头行
  9 空的表不占一节，但 notes 里有交代（几张表、哪张没转出来）
 10 notes 里报的行数与张数与 openpyxl 独立数出来的一致

数字比较用"两边都能读成数就按数值比，否则按字面比"：`1.50` 这种写法本来就是要保住的字面，
而 openpyxl 只会给一个 float —— 按字面硬比会把对的实现判成错。
"""
from __future__ import annotations

import datetime as DT
import os
import re
import sys

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from verify_epub_write import norm  # noqa: E402
from verify_odt_read import (  # noqa: E402
    check,
    pandoc_read,
    plain,
    table_rows,
)

BUILD = os.path.join("core", "build", "xlsxread")
SOURCE = os.path.join(BUILD, "book.xlsx")
KINDS = {"md": "markdown", "html": "html", "docx": "docx"}


def as_text(value) -> str:
    """openpyxl 给的东西变成"该在格子里出现的样子"（日期先按 ISO；数字留给数值比较）。"""
    if value is None:
        return ""
    if isinstance(value, DT.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if (value.hour or value.minute or value.second) else value.date().isoformat()
    if isinstance(value, DT.date):
        return value.isoformat()
    return str(value)


def grid_of(sheet) -> tuple:
    """按**坐标**把这张表摆成网格：行列都到 openpyxl 说的最大行号与最大列号。"""
    widest = sheet.max_column or 0
    rows = []
    for index in range(1, (sheet.max_row or 0) + 1):
        rows.append([as_text(sheet.cell(row=index, column=column).value) for column in range(1, widest + 1)])
    return rows, widest


def same_cell(ours: str, theirs: str) -> bool:
    def number(text):
        try:
            return float(text)
        except (TypeError, ValueError):
            return None

    left, right = number(ours), number(theirs)
    if left is not None and right is not None:
        return left == right
    return norm(ours) == norm(theirs)


def same_grid(a: list, b: list) -> bool:
    if len(a) != len(b):
        return False
    for row_a, row_b in zip(a, b):
        if len(row_a) != len(row_b):
            return False
        for cell_a, cell_b in zip(row_a, row_b):
            if not same_cell(cell_a, cell_b):
                return False
    return True


def headings(ast) -> list:
    """产物里的二级标题文字（pandoc 读回来是 Header 块）。"""
    out = []
    for block in ast["blocks"]:
        if block.get("t") == "Header" and block["c"][0] == 2:
            out.append(norm(re.sub(r"\s+", "", plain(block))))
    return out


def main() -> int:
    if not os.path.isdir(BUILD):
        print("FAIL 0 没有 %s（先跑 :core:test 的落盘那条）" % BUILD)
        return 1
    missing = [name for name in [SOURCE] + [os.path.join(BUILD, "book.%s" % suffix) for suffix in
                                            ("md", "html", "docx", "grid.txt")]
               if not os.path.exists(name) or os.path.getsize(name) == 0]
    check(2, "三条产物与原件都在且非空", not missing, "缺 %s" % missing)
    try:
        book = openpyxl.load_workbook(SOURCE)
    except Exception as error:                      # noqa: BLE001
        check(1, "openpyxl 读得动这份 xlsx", False, "%s: %s" % (type(error).__name__, error))
        return finish()
    check(1, "openpyxl 读得动这份 xlsx", True)
    if missing:
        return finish()

    filled = [(name, grid_of(book[name])[0]) for name in book.sheetnames]
    with_rows = [(name, rows) for name, rows in filled if any(any(cell for cell in row) for row in rows)]
    grids = [(name, rows, grid_of(book[name])[1]) for name, rows in with_rows]

    readings = {}
    for key, fmt in KINDS.items():
        ast = pandoc_read(fmt, os.path.join(BUILD, "book.%s" % key))
        readings[key] = ast
    check(3, "一节一张表：表名与顺序照工作簿声明",
          all(headings(readings[key]) == [norm(name) for name, _, _ in grids] for key in readings),
          "openpyxl %s vs %s" % ([name for name, _, _ in grids], {k: headings(readings[k]) for k in readings}))

    ours = {key: table_rows(readings[key]) for key in readings}
    want = [rows for _, rows, _ in grids]
    off = {key: value for key, value in ours.items() if not same_grid_rows(want, value)}
    check(4, "整张格子逐格对（openpyxl 按坐标摆的那份）", not off,
          "原件 %s vs %s" % (want, {key: off[key] for key in list(off)[:1]}))

    sizes = [(len(rows), width) for _, rows, width in grids]
    read_sizes = {key: [(len(table), max(len(row) for row in table) if table else 0) for table in ours[key]]
                  for key in readings}
    check(5, "网格尺寸：行数与列数都到文件声明的上限（不截尾、不缺列）",
          all(value == sizes for value in read_sizes.values()),
          "应为 %s vs %s" % (sizes, read_sizes))

    def cell_texts(key: str) -> list:
        return [norm(cell) for table in ours[key] for row in table for cell in row]

    timed = [(name, cell.coordinate, cell.value) for name in book.sheetnames
             for row in book[name].iter_rows() for cell in row
             if isinstance(cell.value, DT.datetime) and (cell.value.hour or cell.value.minute or cell.value.second)]
    lost = ["%s!%s 的时间没搬过来（应为 %s）" % (name, coordinate, value.strftime("%H:%M:%S"))
            for name, coordinate, value in timed
            if not any(value.strftime("%H:%M:%S") in text for text in cell_texts("md"))]
    serial = [key for key in readings
              if re.search(r"(?<![\d.])\d{5}(?![\d.])", " ".join(cell_texts(key)))]
    check(6, "日期是 ISO 写法，不是一串序列号（有时间要带时间）", not lost and not serial,
          "%s；看着像序列号的产物：%s" % (lost[:3], serial))

    # 这些值必须**一字不差**地出现在某个格子里：38.5 写成 38.50 就是改了人家的写法
    literals = [norm(word) for word in ("38.5", "0.25", "2023-05-01", "含中文的格子，带逗号,", "纯文字",
                                        "第五行只有 D 列有东西")]
    absent = {key: [word for word in literals if word not in cell_texts(key)] for key in readings}
    check(7, "数字与文字照原样进格子（带逗号的那格也整格留着）",
          all(not value for value in absent.values()),
          str({key: absent[key] for key in absent if absent[key]}))

    heads = {key: [bool((block["c"][3] or [[], []])[1]) for block in readings[key]["blocks"]
                   if block.get("t") == "Table"] for key in readings}
    check(8, "表头那一行留空（文件没说首行是表头就不抬）",
          all(value == [False] * len(grids) for value in heads.values()),
          str(heads))

    notes = open(os.path.join(BUILD, "book.notes.txt"), encoding="utf-8").read() if os.path.exists(
        os.path.join(BUILD, "book.notes.txt")) else ""
    empties = [name for name, rows in filled if not any(any(cell for cell in row) for row in rows)]
    unnamed = [name for name in empties if norm(name) not in norm(notes)]
    check(9, "空的表不占一节，但表名与张数在说明里点出来", not unnamed and "空" in notes,
          "没点名的空表：%s；notes=%s" % (unnamed, notes.replace("\n", " / ")[:160]))

    told = re.search(r"(\d+) 张表 · (\d+) 行", notes)
    per_sheet = re.findall(r"^(.+?)：(\d+) 行 × (\d+) 列$", notes, re.M)
    want_sizes = [(name, len(rows), width) for name, rows, width in grids]
    check(10, "notes 报的张数、行数与每张表的尺寸与 openpyxl 数出来的一致",
          bool(told) and int(told.group(1)) == len(grids)
          and int(told.group(2)) == sum(len(rows) for _, rows, _ in grids)
          and [(name, int(rows), int(width)) for name, rows, width in per_sheet] == want_sizes,
          "notes：%s；openpyxl：%s 张 / %s 行 / %s"
          % (notes.replace("\n", " / ")[:150], len(grids),
             sum(len(rows) for _, rows, _ in grids), want_sizes))
    return finish()


def same_grid_rows(want: list, ours: list) -> bool:
    if len(want) != len(ours):
        return False
    return all(same_grid(a, b) for a, b in zip(want, ours))


def finish() -> int:
    import verify_odt_read

    failed = [item for item in verify_odt_read.results if not item[2]]
    print("\n共 %d 条，%d 条红" % (len(verify_odt_read.results), len(failed)))
    for number, name, _ in failed:
        print("  红的是第 %d 条：%s" % (number, name))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
