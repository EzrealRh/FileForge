#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""拿 openpyxl 复核 Kotlin 读出来的 xlsx。

openpyxl 是一套完整独立的 SpreadsheetML 实现（读写都自己来），所以"同一份文件读出同一批格子"
不是我自己家两个实现互相点头。比对走 **csv 标准库**读 Kotlin 的产物，连引号转义一起验。

跑法：先 `./gradlew :core:test`（产物落到 core/build/office/），再 python tools/verify_xlsx.py
"""
import csv
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = os.path.join(HERE, "..")
PRODUCTS = os.path.join(ROOT, "core", "build", "office")

from make_office_fixtures import OUT, render_cell  # noqa: E402

FIXTURE = os.path.join(OUT, "book.xlsx")

PASSED = []
FAILED = []


def square(rows, width):
    """两边都补齐到同一宽度再比 —— 补齐规矩本身就是要验的东西。"""
    return [row + [""] * (width - len(row)) for row in rows]


def check(what, ok):
    print(("  通过：" if ok else "  失败：") + what)
    (PASSED if ok else FAILED).append(what)
    return ok


def mine(index):
    """Kotlin 那张表的 CSV：用 csv 模块读，不自己按逗号切。"""
    path = os.path.join(PRODUCTS, "book.xlsx.sheet%d.csv" % index)
    if not os.path.isfile(path):
        raise SystemExit("没有 %s，先跑 ./gradlew :core:test" % path)
    with io.open(path, encoding="utf-8", newline="") as handle:
        return [row for row in csv.reader(handle) if row]


def main():
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("本机没有 openpyxl，装一下再跑：pip install openpyxl")
    book = load_workbook(FIXTURE)

    check("张数与表名顺序都跟工作簿声明的一致",
          [s.title for s in book.worksheets] == ["费用", "空表", "第三张"] and
          len(mine(1)) > 0 and mine(2) == [] and mine(3) != [])

    expected = [[render_cell(v) for v in row] for row in book["费用"].iter_rows(values_only=True)]
    expected = [row for row in expected if any(cell for cell in row)]
    got = mine(1)
    width = max([len(row) for row in expected + got])
    check("「费用」每张格子逐字相同（%d 行 × %d 列）" % (len(expected), width),
          square(expected, width) == square(got, width))

    flat = [cell for row in got for cell in row]
    check("日期格子转成了 ISO 而不是序列号 45047",
          "2023-05-01" in flat and not any(cell == "45047" for cell in flat))
    check("带时间的格子保留到秒", "2024-02-29 08:30:15" in flat)
    check("自定义格式（年月日是字面文字）也认成日期", "2021-03-04" in flat)
    check("数字写法没被改写", "38.5" in flat and "0.25" in flat)
    check("带逗号的格子被正确引用，csv 读回来还是原文",
          any("含中文的格子，带逗号," == cell for cell in flat))
    check("稀疏格子按引用对位（第五行只有 D 列）",
          len(got) >= 5 and got[4][3] == "第五行只有 D 列有东西" and not any(got[4][:3]))
    check("尾部整行空白被去掉", got[-1] and any(cell for cell in got[-1]))

    third = [cell for row in mine(3) for cell in row]
    check("第三张表的字落在第三份产物里", third == ["顺序要按工作簿声明的来"])

    print("\n通过 %d 条，失败 %d 条" % (len(PASSED), len(FAILED)))
    if FAILED:
        for item in FAILED:
            print("  未过：" + item)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
